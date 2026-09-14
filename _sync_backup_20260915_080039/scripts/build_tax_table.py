# -*- coding: utf-8 -*-
"""근로소득 간이세액표(국세청 배포 엑셀/CSV) → catalog/tax_table_2026.json 변환기.

매월 떼는 소득세는 계산식이 아니라 소득세법 시행령 별표2의 '표'를 그대로 찾아 쓴다.
그래서 실수령액 계산기는 이 표 원문이 있어야 정확하다.

[사용법]
  1) 홈택스 → 세금신고 → 원천세 신고 → 근로소득 간이세액표 에서 전체 파일(엑셀)을 받아
     catalog/ 에 넣는다.
  2) python scripts\\build_tax_table.py            (catalog 안에서 파일을 자동으로 찾는다)
     python scripts\\build_tax_table.py <파일경로>  (직접 지정)
  3) catalog/tax_table_2026.json 이 만들어지면 build_site.py 가 /tools/salary/ 를 만든다.

표는 매년 2월경 개정된다. 개정본을 받아 같은 자리에 넣고 이 스크립트를 다시 돌리면 된다.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "catalog"
OUT = CATALOG / "tax_table_2026.json"
YEAR = 2026
FAMILY_MAX = 11  # 표의 공제대상가족 수 칸 (1~11명)


class TaxTableError(Exception):
    """표를 읽지 못했을 때."""


def find_source() -> Path:
    """catalog 안에서 간이세액표로 보이는 파일을 찾는다."""
    cands: list[Path] = []
    for pat in ("*.xlsx", "*.xls", "*.csv"):
        cands += [p for p in CATALOG.glob(pat) if "간이세액" in p.name or "tax" in p.name.lower()]
    if not cands:
        cands = [p for p in CATALOG.glob("*.xlsx")] + [p for p in CATALOG.glob("*.csv")]
    if not cands:
        raise TaxTableError(
            "catalog 폴더에서 간이세액표 파일(.xlsx/.csv)을 찾지 못했습니다.\n"
            "홈택스에서 받은 파일을 catalog\\ 에 넣고 다시 실행하세요."
        )
    return max(cands, key=lambda p: p.stat().st_mtime)


def read_rows(path: Path) -> Iterable[list[Any]]:
    """엑셀·CSV를 셀 2차원 배열로 읽는다."""
    if path.suffix.lower() == ".csv":
        import csv
        for enc in ("utf-8-sig", "cp949", "utf-8"):
            try:
                with path.open(encoding=enc, newline="") as fh:
                    return [row for row in csv.reader(fh)]
            except UnicodeDecodeError:
                continue
        raise TaxTableError(f"{path.name} 의 한글 인코딩을 알 수 없습니다.")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 실행 환경에 따라 다름
        raise TaxTableError(
            "엑셀을 읽으려면 openpyxl 이 필요합니다.  pip install openpyxl"
        ) from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    # 국세청 배포본은 시트가 둘이다 — '소득령 별표2'(비고)와 '근로소득간이세액표'(본문).
    # 본문 시트는 행이 압도적으로 많으므로 그것으로 고른다.
    best, best_rows = None, -1
    for name in wb.sheetnames:
        n = wb[name].max_row or 0
        if n > best_rows:
            best, best_rows = name, n
    return [list(r) for r in wb[best].iter_rows(values_only=True)]


def to_num(v: Any) -> float | None:
    """셀 값을 숫자로. '-'(세액 없음)는 0, 숫자가 아니면 None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in {"-", "‐", "–", "—"}:
        return 0.0
    s = re.sub(r"천원.*$", "", s)          # '10,000천원' 같은 마지막 행
    s = re.sub(r"[,\s원]", "", s)
    if not s or not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    return float(s)


def parse(cells: Iterable[list[Any]]) -> dict[str, Any]:
    """표 본문만 골라 구간 하한(원)과 가족수별 세액을 뽑는다.

    본문 행 모양: 월급여액(이상) · 월급여액(미만) · 공제대상가족 1~11명 세액.
    월급여액은 천원 단위이고, 세액이 없는 칸은 '-' 로 적혀 있다(0원).
    구간 간격은 낮은 구간 5천원 → 높은 구간에서 넓어지므로 하한 목록을 그대로 싣는다.
    마지막 행은 '10,000천원' 한 칸짜리다.
    """
    bounds: list[int] = []
    rows: list[list[int]] = []

    for row in cells:
        if len(row) < FAMILY_MAX + 2:
            continue
        lo = to_num(row[0])
        if lo is None or lo <= 0:
            continue
        taxes_raw = [to_num(c) for c in row[2:2 + FAMILY_MAX]]
        if any(t is None for t in taxes_raw):
            continue                       # 머리글·주석 행
        taxes = [int(round(t)) for t in taxes_raw]      # type: ignore[arg-type]
        if any(t < 0 for t in taxes):
            continue
        # 세액은 가족 수가 늘수록 줄어든다. 이 성질로 남은 잡음 행을 거른다.
        if any(taxes[i] < taxes[i + 1] for i in range(FAMILY_MAX - 1)):
            continue
        bounds.append(int(round(lo * 1000)))
        rows.append(taxes)

    if len(rows) < 100:
        raise TaxTableError(
            f"표 본문을 {len(rows)}행밖에 읽지 못했습니다. 홈택스에서 받은 "
            "'근로소득 간이세액표' 전체 파일이 맞는지 확인하세요."
        )

    pair = sorted(zip(bounds, rows), key=lambda x: x[0])
    bounds, rows = [], []
    for b, r in pair:
        if bounds and b == bounds[-1]:
            rows[-1] = r
            continue
        bounds.append(b)
        rows.append(r)

    steps = sorted({bounds[i + 1] - bounds[i] for i in range(len(bounds) - 1)})
    return {
        "year": YEAR,
        "family_max": FAMILY_MAX,
        "min": bounds[0],
        "max": bounds[-1],
        "step": steps[0] if steps else 5000,
        "uniform": len(steps) == 1,
        "bounds": bounds,
        "rows": rows,
        "source": "소득세법 시행령 별표2 근로소득 간이세액표 (국세청 배포본)",
    }


def main(argv: list[str]) -> int:
    try:
        src = Path(argv[1]) if len(argv) > 1 else find_source()
        if not src.exists():
            raise TaxTableError(f"{src} 를 찾을 수 없습니다.")
        table = parse(read_rows(src))
    except TaxTableError as exc:
        print(f"[간이세액표] 실패: {exc}")
        return 1

    OUT.write_text(json.dumps(table, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    size = OUT.stat().st_size / 1024
    print(f"[간이세액표] {src.name} → {OUT.name}")
    print(f"[간이세액표] 구간 {len(table['rows']):,}개 · 월급여 "
          f"{table['min']:,}원 ~ {table['max']:,}원 · 간격 {table['step']:,}원 "
          f"({'일정' if table['uniform'] else '구간별로 다름'}) · {size:,.0f} KB")
    mid = len(table["rows"]) // 2
    print(f"[간이세액표] 표본: 월급여 {table['bounds'][mid]:,}원 · 가족 1명 "
          f"{table['rows'][mid][0]:,}원 / 2명 {table['rows'][mid][1]:,}원")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

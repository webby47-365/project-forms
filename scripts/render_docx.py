"""명세(FormSpec) → DOCX 마스터 생성 (python-docx).

DOCX가 편집용 원본이며, 여기서 PDF를 파생한다.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from form_spec import FONT_SANS, MARGIN_MM, USABLE_WIDTH_MM, FormSpec, pct_to_mm

# ── 디자인 토큰 (현대적 서식 스타일) ──
ACCENT = "24384F"        # 강조색: 딥 슬레이트 네이비 (표 머리글 배경, 강조선)
HEADER_FILL = ACCENT     # 표 머리글 배경
HEADER_TEXT = "FFFFFF"   # 표 머리글 글자
LABEL_FILL = "F4F6F9"    # 라벨 셀 배경 (아주 연한 회청)
LABEL_TEXT = "2B3A4A"    # 라벨 글자
BORDER_COLOR = "C9D2DD"  # 본문 테두리 (얇고 연하게)
MUTED = "6B7280"         # 보조 문구
SAMPLE_TEXT = "1F5C8C"   # 작성 예시 글자색 (라벨·본문과 구분되는 슬레이트 블루)
SITE_NAME = "무료서식 다운로드"  # 문서 하단 출처 표기 (빈 문자열이면 표기 생략)

# 증명사진 자리에 넣을 실루엣 이미지 (예시 모드 전용). 없으면 사진칸은 글자만 남는다.
PHOTO_SAMPLE = Path(__file__).resolve().parent.parent / "assets" / "photo-sample.png"
PHOTO_DEFAULT_MM = (26.0, 34.0)  # 3×4cm 증명사진 비율

# 자동 맞춤 배율: 행 높이와 문단 간격에만 적용되며 글자 크기는 건드리지 않는다.
_SCALE = 1.0
# 예시 모드: 미리보기 이미지 전용. 다운로드 파일은 항상 빈 양식으로 렌더링한다.
_SAMPLE = False


def _set_font(run: Any, size_pt: float, *, bold: bool = False, color: str | None = None) -> None:
    """런의 한글/영문 폰트를 함께 지정한다 (eastAsia 미지정 시 한글이 깨짐)."""
    run.font.name = FONT_SANS
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_SANS)


def _para(
    doc: Any,
    text: str = "",
    *,
    size: float = 10.5,
    bold: bool = False,
    align: str = "left",
    space_after: float = 4.0,
    space_before: float = 0.0,
    color: str | None = None,
    indent_mm: float = 0.0,
    prefix: tuple[str, str] | None = None,
    rule: str | None = None,
    rule_size: int = 12,
) -> Any:
    """문단 추가.

    Args:
        prefix: (텍스트, 색상) 형태의 강조 접두 런. 소제목 앞 악센트 바 등에 사용.
        rule: 지정 시 문단 하단에 해당 색상의 강조선을 그린다.
    """
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_after = Pt(space_after * _SCALE)
    pf.space_before = Pt(space_before * _SCALE)
    pf.line_spacing = 1.3
    if indent_mm:
        pf.left_indent = Mm(indent_mm)
    if prefix is not None:
        pre = p.add_run(prefix[0])
        _set_font(pre, size, bold=True, color=prefix[1])
    run = p.add_run(text)
    _set_font(run, size, bold=bold, color=color)
    if rule is not None:
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(rule_size))
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), rule)
        borders.append(bottom)
        p_pr.append(borders)
    return p


def _shade_cell(cell: Any, fill: str) -> None:
    """셀 배경 음영."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_borders(cell: Any) -> None:
    """셀 4면 실선 테두리."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER_COLOR)
        borders.append(el)
    tc_pr.append(borders)


def _cell_text(
    cell: Any,
    text: str,
    *,
    size: float = 10.0,
    bold: bool = False,
    align: str = "left",
    fill: str | None = None,
    color: str | None = None,
) -> None:
    """셀에 텍스트를 넣고 서식을 적용한다 (여러 줄은 \\n으로 분리)."""
    cell.text = ""
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[align]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.3
        run = p.add_run(line)
        _set_font(run, size, bold=bold, color=color)
    # tcPr 자식 순서는 스키마로 강제됨: tcBorders → shd → vAlign
    _cell_borders(cell)
    if fill:
        _shade_cell(cell, fill)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    cell._tc.get_or_add_tcPr().append(v)


def _row_height(table: Any, row_idx: int, min_height_mm: float, *, absolute: bool = False) -> None:
    """행 최소 높이를 1회만 설정한다 (셀마다 설정하면 trHeight가 중복되어 높이가 틀어짐).

    Args:
        absolute: True면 자동 맞춤 배율(_SCALE)을 곱하지 않은 실제 mm로 설정한다(사진칸 보정용).
    """
    tr_pr = table.rows[row_idx]._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(existing)
    h = OxmlElement("w:trHeight")
    mm = min_height_mm if absolute else min_height_mm * _SCALE
    h.set(qn("w:val"), str(math.ceil(mm * 56.7) if absolute else int(mm * 56.7)))  # mm → twip
    h.set(qn("w:hRule"), "atLeast")
    tr_pr.append(h)


def _new_table(doc: Any, rows: int, cols: int, widths_mm: list[float]) -> Any:
    """표 생성 후 컬럼 폭 고정 (자동맞춤 해제)."""
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Mm(widths_mm[idx])
    return table


def _sample_cell(blk: dict[str, Any], r: int, c: int) -> str:
    """예시 모드에서 (r, c) 칸에 넣을 예시 값. 없으면 빈 문자열."""
    if not _SAMPLE:
        return ""
    rows = blk.get("sample_rows")
    if not rows or r >= len(rows):
        return ""
    row = rows[r]
    if not isinstance(row, list) or c >= len(row):
        return ""
    return str(row[c] or "")


PHOTO_PAD_H_MM = 2.0   # 사진 위아래로 남길 여유 (셀 안여백 + 선 두께)
# 사진 좌우로 남길 여유(격자 열 폭 기준). LibreOffice는 안여백을 0으로 둔 셀(_center_cell)에서도
# 내용을 왼쪽 선에서 약 2.9mm 띄워 그린다. 사진이 그 안쪽 폭보다 넓으면 가운데로 오지 않고 왼쪽에 붙어
# 오른쪽이 칸 밖으로 잘린다. 24.85mm 격자 칸에서 폭별로 재어 보니 18.9mm(여유 6mm)부터 가운데에 놓였다.
PHOTO_PAD_W_MM = 6.0


def _photo_span(blk: dict[str, Any], pr: int, pc: int, n_rows: int, n_cols: int) -> tuple[int, int]:
    """사진칸이 병합된 (행 수, 열 수). 병합이 없으면 (1, 1)."""
    for merge in blk.get("merges", []):
        r1, c1, r2, c2 = merge
        if r1 == pr and c1 == pc and 0 <= r2 < n_rows and 0 <= c2 < n_cols:
            return r2 - r1 + 1, c2 - c1 + 1
    return 1, 1


def _photo_size(w_mm: float, h_mm: float, pc: int, span_cols: int,
                widths_mm: list[float], n_cols: int) -> tuple[float, float]:
    """사진 크기(mm). 명세의 photo_mm를 1페이지 자동 맞춤 배율과 무관하게 그대로 쓴다.

    예전에는 칸 높이·배율에 맞춰 사진을 줄여 이력서마다 실루엣 크기와 위치가 제각각이었다.
    이제 높이가 모자라면 사진이 아니라 행을 늘리고(_photo_rows), 열 폭이 모자랄 때만
    폭에 맞춰 비율을 유지한 채 줄인다.

    열 폭은 명세 widths 와 실제 표 격자 중 좁은 쪽을 쓴다. python-docx 는 표 격자(w:tblGrid)를
    본문 폭 ÷ 열 수로 똑같이 나눠 만들고, 고정 레이아웃에서 LibreOffice 는 이 격자로 그린다
    (이력서: 명세 31.3mm, 실제 약 24.9mm). 명세 폭만 보면 사진이 칸을 넘어 오른쪽이 잘린다.
    """
    grid_w = USABLE_WIDTH_MM * span_cols / n_cols
    avail_w = min(sum(widths_mm[pc:pc + span_cols]), grid_w) - PHOTO_PAD_W_MM
    if w_mm > 0 and avail_w < w_mm:
        ratio = avail_w / w_mm
        return round(w_mm * ratio, 1), round(h_mm * ratio, 1)
    return w_mm, h_mm


def _photo_rows(table: Any, blk: dict[str, Any], pr: int, span_rows: int, photo_h_mm: float) -> None:
    """사진칸 병합 행들의 높이 합이 (사진 높이 + 여유)보다 작으면 행 최소 높이를 균등하게 올린다.

    빈 양식에도 적용해 3×4cm 사진을 붙일 자리를 확보한다. 배율로 행이 충분히 크면 그대로 둔다.
    """
    row_h = float(blk.get("row_height_mm", 7.5)) * _SCALE
    need = photo_h_mm + PHOTO_PAD_H_MM
    if row_h * span_rows >= need:
        return
    for r in range(pr, pr + span_rows):
        _row_height(table, r, need / span_rows, absolute=True)


def _center_cell(cell: Any) -> None:
    """사진칸 셀: 좌우 안여백 0 + 세로 가운데 정렬 (tcMar·vAlign 을 하나씩만 남긴다).

    tcPr 자식 순서는 스키마로 정해져 있어 tcMar 를 vAlign 앞에 둔다.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcMar")) + tc_pr.findall(qn("w:vAlign")):
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side in ("left", "right"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), "0")
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tc_pr.append(mar)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), "center")
    tc_pr.append(v)


def _insert_photo(cell: Any, width_mm: float, height_mm: float) -> bool:
    """셀에 증명사진 실루엣을 가로·세로 가운데로 넣는다. 이미지 파일이 없으면 False."""
    if not PHOTO_SAMPLE.exists():
        return False
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    # 줄 간격 배수(1.3)가 그림 높이에도 곱해지면 사진이 칸보다 커져 밀려나므로 한 줄 간격으로 둔다
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(PHOTO_SAMPLE), width=Mm(width_mm), height=Mm(height_mm))
    _center_cell(cell)
    return True


def render_docx(spec: FormSpec, out_path: Path, *, scale: float = 1.0,
                sample: bool = False) -> Path:
    """명세를 DOCX로 렌더링한다.

    Args:
        scale: 행 높이·문단 간격 축소 배율. 1페이지 수납 자동 맞춤에 사용한다.
        sample: True면 명세의 예시값(sample_rows·sample_text 등)을 채워 넣는다.
            미리보기 이미지 전용이며, 사용자가 내려받는 파일에는 절대 쓰지 않는다.
            레이아웃(행 수·행 높이·폭)은 빈 양식과 완전히 동일하게 유지된다.
    """
    global _SCALE, _SAMPLE
    _SCALE = scale
    _SAMPLE = sample
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Mm(MARGIN_MM))

    style = doc.styles["Normal"]
    style.font.name = FONT_SANS
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SANS)

    if SITE_NAME:
        footer_p = sec.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_p.paragraph_format.space_before = Pt(0)
        footer_p.paragraph_format.space_after = Pt(0)
        _set_font(footer_p.add_run(SITE_NAME), 7.5, color=MUTED)

    for blk in spec.blocks:
        _render_block(doc, blk)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _render_block(doc: Any, blk: dict[str, Any]) -> None:
    """블록 1개를 렌더링한다."""
    btype = blk["type"]

    if btype == "doc_title":
        # 제목은 글자 사이를 벌려 서식 문서 느낌을 준다
        text = str(blk["text"])
        spaced = blk.get("spaced", True)
        shown = "  ".join(text) if spaced and len(text) <= 8 else text
        _para(doc, shown, size=18, bold=True, align="center",
              space_after=2, space_before=0, color=ACCENT)
        _para(doc, "", size=2, space_after=5, rule=ACCENT, rule_size=14)

    elif btype == "subtitle":
        _para(doc, str(blk["text"]), size=10.5, bold=True, align=blk.get("align", "left"),
              space_before=4, space_after=2, color=LABEL_TEXT,
              prefix=("▌ ", ACCENT), rule=BORDER_COLOR, rule_size=6)

    elif btype == "para":
        text = str(blk.get("text", ""))
        if _SAMPLE and blk.get("sample_text"):
            text = str(blk["sample_text"])
        _para(doc, text, size=blk.get("size", 10.5),
              align=blk.get("align", "left"), bold=blk.get("bold", False),
              space_after=blk.get("space_after", 4), indent_mm=blk.get("indent_mm", 0),
              color=SAMPLE_TEXT if (_SAMPLE and blk.get("sample_text")) else None)

    elif btype == "grid":
        rows: list[list[str]] = blk["rows"]
        widths = pct_to_mm(blk.get("widths") or [100 / len(rows[0])] * len(rows[0]))
        label_cols = set(blk.get("label_cols", [i for i in range(len(rows[0])) if i % 2 == 0]))
        photo_cell = blk.get("photo_cell")
        table = _new_table(doc, len(rows), len(rows[0]), widths)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                is_label = c in label_cols and str(val).strip() != ""
                sample = _sample_cell(blk, r, c)
                # 라벨 칸은 예시 모드에서도 그대로 두고, 빈 입력칸에만 예시값을 채운다
                if sample and not is_label:
                    _cell_text(table.cell(r, c), sample, size=10.0, align="left",
                               color=SAMPLE_TEXT)
                else:
                    _cell_text(
                        table.cell(r, c), str(val), size=10.0, bold=is_label,
                        align="center" if is_label else "left",
                        fill=LABEL_FILL if is_label else None,
                        color=LABEL_TEXT if is_label else None,
                    )
            _row_height(table, r, blk.get("row_height_mm", 7.5))
        n_rows, n_cols = len(rows), len(rows[0])
        for merge in blk.get("merges", []):
            r1, c1, r2, c2 = merge
            if not (0 <= r1 <= r2 < n_rows and 0 <= c1 <= c2 < n_cols):
                print(f"  [경고] merges {merge}가 표({n_rows}x{n_cols}) 범위를 벗어나 무시했습니다.")
                continue
            merged = table.cell(r1, c1).merge(table.cell(r2, c2))
            # 병합 시 각 셀의 빈 문단이 누적되어 행 높이가 비정상적으로 커지므로 정리한다
            keep = merged.paragraphs[0]
            for extra in merged.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)
            if not keep.runs:
                _set_font(keep.add_run(""), 10.0)
        # 증명사진 칸: 사진은 명세 크기(photo_mm) 그대로. 칸 높이가 모자라면 행을 늘린다.
        # 행 높이 보정은 빈 양식에도 적용해 사진 붙일 자리를 확보하고,
        # 실루엣 이미지는 예시 모드(미리보기)에서만 넣는다.
        if photo_cell:
            pr, pc = int(photo_cell[0]), int(photo_cell[1])
            if 0 <= pr < n_rows and 0 <= pc < n_cols:
                span_rows, span_cols = _photo_span(blk, pr, pc, n_rows, n_cols)
                w_mm, h_mm = blk.get("photo_mm", PHOTO_DEFAULT_MM)
                w_mm, h_mm = _photo_size(float(w_mm), float(h_mm), pc, span_cols, widths, n_cols)
                _photo_rows(table, blk, pr, span_rows, h_mm)
                cell = table.cell(pr, pc)
                _center_cell(cell)
                if _SAMPLE:
                    _insert_photo(cell, w_mm, h_mm)
        _para(doc, "", size=2, space_after=blk.get("space_after", 2))

    elif btype == "table":
        header: list[str] = blk["header"]
        widths = pct_to_mm(blk.get("widths") or [100 / len(header)] * len(header))
        n_empty = int(blk.get("empty_rows", 3))
        body: list[list[str]] = blk.get("rows", [])
        # 예시 모드: 빈 입력행을 앞에서부터 예시값으로 채운다. 행 수는 그대로 유지한다.
        samples: list[list[str]] = list(blk.get("sample_rows") or []) if _SAMPLE else []
        samples = samples[:n_empty]
        table = _new_table(doc, 1 + len(body) + n_empty, len(header), widths)
        for c, h in enumerate(header):
            _cell_text(table.cell(0, c), h, size=9.5, bold=True, align="center",
                       fill=HEADER_FILL, color=HEADER_TEXT)
        _row_height(table, 0, blk.get("row_height_mm", 7.5))
        for r, row in enumerate(body, start=1):
            for c, val in enumerate(row):
                _cell_text(table.cell(r, c), str(val), size=10.0,
                           align=blk.get("body_align", "center"))
            _row_height(table, r, blk.get("row_height_mm", 7.5))
        base = 1 + len(body)
        for i in range(n_empty):
            r = base + i
            srow = samples[i] if i < len(samples) else []
            for c in range(len(header)):
                val = str(srow[c]) if c < len(srow) and srow[c] is not None else ""
                _cell_text(table.cell(r, c), val, size=10.0,
                           align=blk.get("body_align", "center"),
                           color=SAMPLE_TEXT if val else None)
            _row_height(table, r, blk.get("row_height_mm", 7.5))
        _para(doc, "", size=2, space_after=blk.get("space_after", 2))

    elif btype == "textbox":
        # 자유기술란: 안내문은 작은 회색 글씨로 박스 위에 두고, 박스는 빈 칸으로 남긴다
        hint = str(blk.get("hint", "")).strip()
        if hint:
            _para(doc, hint, size=8.5, color=MUTED, space_after=2)
        widths = pct_to_mm([100])
        table = _new_table(doc, 1, 1, widths)
        sample_text = str(blk.get("sample_text", "")).strip() if _SAMPLE else ""
        _cell_text(table.cell(0, 0), sample_text, size=9.5, align="left",
                   color=SAMPLE_TEXT if sample_text else None)
        if sample_text:
            # 예시 글이 위에서부터 채워지도록 세로 정렬을 위쪽으로 바꾼다
            tc_pr = table.cell(0, 0)._tc.get_or_add_tcPr()
            for v in tc_pr.findall(qn("w:vAlign")):
                v.set(qn("w:val"), "top")
        _row_height(table, 0, blk.get("height_mm", 30))
        _para(doc, "", size=2, space_after=blk.get("space_after", 2))

    elif btype == "kv_list":
        for item in blk["items"]:
            _para(doc, f"· {item}", size=10.0, space_after=2, indent_mm=3)

    elif btype == "article":
        _para(doc, str(blk["heading"]), size=10.5, bold=True, space_before=5, space_after=2)
        for line in blk.get("lines", []):
            _para(doc, str(line), size=10.0, align="both", space_after=2, indent_mm=3)

    elif btype == "date_line":
        if _SAMPLE and blk.get("sample_text"):
            _para(doc, str(blk["sample_text"]), size=11, align="center",
                  space_before=4, space_after=4, color=SAMPLE_TEXT)
        else:
            _para(doc, blk.get("text", "20        년        월        일"),
                  size=11, align="center", space_before=4, space_after=4)

    elif btype == "sign_line":
        labels = blk.get("labels", [blk.get("label", "작성자")])
        names = list(blk.get("sample_names") or []) if _SAMPLE else []
        for i, label in enumerate(labels):
            name = str(names[i]) if i < len(names) and names[i] else ""
            if name:
                p = _para(doc, "", size=11, align="right", space_after=4)
                _set_font(p.runs[0] if p.runs else p.add_run(""), 11)
                run_label = p.add_run(f"{label} : ")
                _set_font(run_label, 11)
                run_name = p.add_run(f"{name}          ")
                _set_font(run_name, 11, color=SAMPLE_TEXT)
                run_tail = p.add_run("(서명 또는 인)")
                _set_font(run_tail, 11)
            else:
                _para(doc, f"{label} :                                        (서명 또는 인)",
                      size=11, align="right", space_after=4)

    elif btype == "notice":
        _para(doc, "※ " + str(blk["text"]), size=8.5, align="left",
              space_before=3, space_after=0, color=MUTED)

    elif btype == "spacer":
        _para(doc, "", size=blk.get("size", 10), space_after=blk.get("height_pt", 8))

    elif btype == "page_break":
        doc.add_page_break()
